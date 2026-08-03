"""Build a read-only institutional due-diligence pack from sealed evidence.

This exporter does not fit a model, select a rule, recompute predictions or open
a statistical trial. It packages existing registered artifacts and verifies the
resulting workbook/CSVs against an input/output SHA-256 manifest.
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "reports/usdcop_directional_institutional_pack_v1"
WORKBOOK = OUTPUT_DIR / "usdcop_directional_due_diligence_v1.xlsx"
MANIFEST = OUTPUT_DIR / "evidence_manifest.json"

INPUTS = {
    "acceptance_standard": "config/forecast_experiments/institutional_directional_acceptance_v1.yaml",
    "weekly_contract": "config/forecast_experiments/usdcop_h1_regime_shadow_v2.yaml",
    "weekly_registration": "reports/usdcop_h1_regime_shadow_v2/registration.json",
    "weekly_evaluation": "reports/usdcop_h1_regime_shadow_v2/prospective_evaluation.json",
    "daily_contract": "config/forecast_experiments/usdcop_h1_daily_shadow_v1.yaml",
    "daily_registration": "reports/usdcop_h1_daily_shadow_v1/registration.json",
    "daily_evaluation": "reports/usdcop_h1_daily_shadow_v1/prospective_evaluation.json",
    "daily_public_index": "usdcop-trading-dashboard/public/forecasting/usdcop/h1_daily_shadow_index.json",
    "regime_research_manifest": "reports/usdcop_causal_regime_gate_manifest.json",
    "forward_flow_levels": "reports/usdcop_forward_flow_direction_v1/manifest.json",
    "forward_flow_stationary": "reports/usdcop_forward_flow_stationary_v2/manifest.json",
    "intraday_latam": "reports/usdcop_intraday_latam_lead_v1/manifest.json",
    "transport_contract": "config/forecast_experiments/usdcop_h1_latam_transport_v1.yaml",
    "transport_registration": "config/forecast_experiments/preregistrations/usdcop_h1_latam_transport_v1.json",
    "transport_result": "reports/usdcop_h1_latam_transport_v1/result.json",
    "transport_metrics": "reports/usdcop_h1_latam_transport_v1/metrics.csv",
    "transport_annual": "reports/usdcop_h1_latam_transport_v1/annual_metrics.csv",
    "directional_trial_registry": ".claude/specs/assets/usdcop/EXP-DIR-001-directional-trials.md",
    "hypothesis_registry": ".claude/specs/assets/usdcop/HYPOTHESIS-REGISTRY.md",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(key: str) -> dict[str, Any]:
    return json.loads((ROOT / INPUTS[key]).read_text(encoding="utf-8"))


def read_yaml(key: str) -> dict[str, Any]:
    return yaml.safe_load((ROOT / INPUTS[key]).read_text(encoding="utf-8"))


def percent(value: Any) -> float | None:
    return None if value is None else float(value)


def evidence_frames() -> tuple[dict[str, pd.DataFrame], dict[str, Any]]:
    acceptance = read_yaml("acceptance_standard")
    weekly_contract = read_yaml("weekly_contract")
    weekly_eval = read_json("weekly_evaluation")
    daily_contract = read_yaml("daily_contract")
    daily_eval = read_json("daily_evaluation")
    regime_manifest = read_json("regime_research_manifest")
    flow_levels = read_json("forward_flow_levels")
    flow_stationary = read_json("forward_flow_stationary")
    intraday = read_json("intraday_latam")
    transport = read_json("transport_result")
    transport_registration = read_json("transport_registration")
    historical = weekly_contract["pre_registration_evidence"]
    audit = historical["audit_2020_2024"]

    claim_ladder = pd.DataFrame([
        {
            "claim_level": "research_candidate",
            "current_status": "PASS",
            "permitted_language": acceptance["claim_levels"]["research_candidate"]["permitted_language"],
            "capital_authorized": False,
            "evidence": "Reproducible historical artifacts, full trial count and limitations disclosed.",
        },
        {
            "claim_level": "prospective_review_eligible",
            "current_status": "NOT_YET",
            "permitted_language": acceptance["claim_levels"]["prospective_review_eligible"]["permitted_language"],
            "capital_authorized": False,
            "evidence": "Daily 0/252 matured base, 0/100 selective, 0/12 months; weekly 0/104 committed.",
        },
        {
            "claim_level": "strong_directional_evidence",
            "current_status": "NOT_REACHED",
            "permitted_language": acceptance["claim_levels"]["strong_directional_evidence"]["permitted_language"],
            "capital_authorized": False,
            "evidence": "No prospective outcomes; historical H1 lift CI crosses zero; external transport failed.",
        },
        {
            "claim_level": "execution_validated",
            "current_status": "NOT_REACHED",
            "permitted_language": acceptance["claim_levels"]["execution_validated"]["permitted_language"],
            "capital_authorized": False,
            "evidence": "No preregistered executable bid/ask, fill, slippage, financing and capacity evidence.",
        },
        {
            "claim_level": "institutional_capital_candidate",
            "current_status": "NOT_REACHED",
            "permitted_language": acceptance["claim_levels"]["institutional_capital_candidate"]["permitted_language"],
            "capital_authorized": False,
            "evidence": "Independent validation, execution validation and governed allocation approval remain absent.",
        },
    ])

    model_status = pd.DataFrame([
        {
            "experiment": weekly_contract["_meta"]["experiment_id"],
            "role": "primary weekly prospective shadow",
            "frequency": "weekly Friday origin",
            "first_eligible": weekly_contract["prospective_protocol"]["first_eligible_iso_week"],
            "status": weekly_eval["status"],
            "committed": weekly_eval["sample"]["committed_weeks"],
            "matured_primary": weekly_eval["sample"]["matured_signals"],
            "minimum_primary": weekly_eval["sample"]["minimum_matured_signals"],
            "registration_valid": weekly_eval["integrity"]["registration_valid"],
            "signal_authorized": weekly_eval["signal_authorized"],
            "capital_authorized": weekly_eval["capital_authorized"],
        },
        {
            "experiment": daily_contract["_meta"]["experiment_id"],
            "role": "faster prospective all-origin shadow",
            "frequency": "daily; model fit once per ISO week",
            "first_eligible": str(daily_contract["prospective_protocol"]["first_eligible_date"]),
            "status": daily_eval["status"],
            "committed": daily_eval["sample"]["committed_sessions"],
            "matured_primary": daily_eval["sample"]["matured_base_predictions"],
            "minimum_primary": daily_contract["prospective_protocol"]["minimum_matured_base_predictions"],
            "registration_valid": daily_eval["integrity"]["registration_valid"],
            "signal_authorized": daily_eval["signal_authorized"],
            "capital_authorized": daily_eval["capital_authorized"],
        },
        {
            "experiment": transport["experiment_id"],
            "role": "locked external historical transport falsification",
            "frequency": "weekly USD/MXN + USD/BRL",
            "first_eligible": "2020-01-01 primary start",
            "status": transport["verdict"],
            "committed": transport["primary_combined"]["n_total"],
            "matured_primary": transport["primary_combined"]["n_signals"],
            "minimum_primary": 180,
            "registration_valid": True,
            "signal_authorized": transport["signal_authorized"],
            "capital_authorized": transport["capital_authorized"],
        },
    ])

    research_metrics = pd.DataFrame([
        {
            "evidence_slice": "USD/COP H1 regime-gated 2020-2024",
            "evidence_class": historical["evidence_classification"],
            "n_signals": audit["signals"],
            "coverage": percent(audit["coverage"]),
            "directional_accuracy": percent(audit["directional_accuracy"]),
            "balanced_accuracy": percent(audit["balanced_accuracy"]),
            "causal_majority_accuracy": None,
            "lift": percent(audit["lift_vs_causal_majority"]),
            "lift_ci_low": percent(audit["lift_block_bootstrap_ci95"][0]),
            "lift_ci_high": percent(audit["lift_block_bootstrap_ci95"][1]),
            "p_value": percent(audit["lift_block_bootstrap_one_sided_p"]),
            "primary_for_claim": False,
            "interpretation": "Promising exploratory candidate; inference misses the fixed 5% gate and is annually unstable.",
        },
        {
            "evidence_slice": "USD/COP H1 regime-gated 2025 audit",
            "evidence_class": "historical audit already inspected",
            "n_signals": historical["audit_2025"]["signals"],
            "coverage": None,
            "directional_accuracy": percent(historical["audit_2025"]["directional_accuracy"]),
            "balanced_accuracy": None,
            "causal_majority_accuracy": None,
            "lift": None,
            "lift_ci_low": None,
            "lift_ci_high": None,
            "p_value": None,
            "primary_for_claim": False,
            "interpretation": "Only 11 signals; not a fresh prospective holdout.",
        },
        {
            "evidence_slice": "USD/COP H1 regime-gated 2026 through W30",
            "evidence_class": "historical audit already inspected",
            "n_signals": historical["audit_2026_through_w30"]["signals"],
            "coverage": None,
            "directional_accuracy": percent(historical["audit_2026_through_w30"]["directional_accuracy"]),
            "balanced_accuracy": None,
            "causal_majority_accuracy": None,
            "lift": None,
            "lift_ci_low": None,
            "lift_ci_high": None,
            "p_value": None,
            "primary_for_claim": False,
            "interpretation": "Only 5 signals and pre-registration; descriptive only.",
        },
        {
            "evidence_slice": "External LatAm transport 2020-2025 combined",
            "evidence_class": "locked reconstructed historical transport",
            "n_signals": transport["primary_combined"]["n_signals"],
            "coverage": percent(transport["primary_combined"]["coverage"]),
            "directional_accuracy": percent(transport["primary_combined"]["directional_accuracy"]),
            "balanced_accuracy": percent(transport["primary_combined"]["balanced_accuracy"]),
            "causal_majority_accuracy": percent(transport["primary_combined"]["causal_majority_accuracy"]),
            "lift": percent(transport["primary_combined"]["lift_vs_causal_majority"]),
            "lift_ci_low": percent(transport["paired_week_cluster_bootstrap"]["ci_low"]),
            "lift_ci_high": percent(transport["paired_week_cluster_bootstrap"]["ci_high"]),
            "p_value": percent(transport["paired_week_cluster_bootstrap"]["one_sided_p"]),
            "primary_for_claim": True,
            "interpretation": "Failed seven gates; closes cross-LatAm transport family without reformulation.",
        },
        {
            "evidence_slice": "External LatAm transport 2026 diagnostic",
            "evidence_class": "non-primary diagnostic",
            "n_signals": transport["diagnostic_2026"]["combined"]["n_signals"],
            "coverage": percent(transport["diagnostic_2026"]["combined"]["coverage"]),
            "directional_accuracy": percent(transport["diagnostic_2026"]["combined"]["directional_accuracy"]),
            "balanced_accuracy": percent(transport["diagnostic_2026"]["combined"]["balanced_accuracy"]),
            "causal_majority_accuracy": percent(transport["diagnostic_2026"]["combined"]["causal_majority_accuracy"]),
            "lift": percent(transport["diagnostic_2026"]["combined"]["lift_vs_causal_majority"]),
            "lift_ci_low": None,
            "lift_ci_high": None,
            "p_value": percent(transport["diagnostic_2026"]["combined"]["pesaran_timmermann_p"]),
            "primary_for_claim": False,
            "interpretation": "N=16; cannot rescue primary failure and includes severe class imbalance.",
        },
    ])

    failed_families = pd.DataFrame([
        {
            "family": "causal regime generalization H1/H5/H10/H15/H20/H25/H30",
            "tests_or_horizons": len(regime_manifest["generalization"]),
            "primary_pass_count": sum(row["passes_all_research_metric_gates"] for row in regime_manifest["generalization"]),
            "evidence_class": "retrospective research",
            "prospective_candidates": 1,
            "capital_authorized": False,
            "disposition": "Only H1 frozen for prospective adjudication; no historical horizon passed all gates.",
        },
        {
            "family": flow_levels["experiment_id"],
            "tests_or_horizons": len(flow_levels["horizons_tested"]),
            "primary_pass_count": flow_levels["primary_gate_pass_count"],
            "evidence_class": flow_levels["evidence_class"],
            "prospective_candidates": flow_levels["prospective_shadow_candidate_count"],
            "capital_authorized": flow_levels["capital_authorized"],
            "disposition": "Closed representation; consolidated historical flow is reconstructed, not true PIT.",
        },
        {
            "family": flow_stationary["experiment_id"],
            "tests_or_horizons": len(flow_stationary["horizons_tested"]),
            "primary_pass_count": flow_stationary["primary_gate_pass_count"],
            "evidence_class": flow_stationary["evidence_class"],
            "prospective_candidates": flow_stationary["prospective_shadow_candidate_count"],
            "capital_authorized": flow_stationary["capital_authorized"],
            "disposition": "Closed after one stationary reformulation; no horizon passed.",
        },
        {
            "family": intraday["experiment_id"],
            "tests_or_horizons": len(intraday["horizons_tested"]),
            "primary_pass_count": intraday["primary_gate_pass_count"],
            "evidence_class": intraday["evidence_class"],
            "prospective_candidates": intraday["prospective_shadow_candidate_count"],
            "capital_authorized": intraday["capital_authorized"],
            "disposition": "Closed; H1 retrospective lift did not survive inference or 2025/2026 audits.",
        },
        {
            "family": transport["experiment_id"],
            "tests_or_horizons": transport_registration["new_primary_trials"],
            "primary_pass_count": int(transport["primary_pass"]),
            "evidence_class": "locked reconstructed external transport",
            "prospective_candidates": 0,
            "capital_authorized": transport["capital_authorized"],
            "disposition": transport["verdict"],
        },
    ])

    gaps = pd.DataFrame([
        {"gate": "daily elapsed calendar months", "current": daily_eval["sample"]["elapsed_days"] / 30.4375, "required": 12, "status": "FAIL_NOT_MATURE", "earliest_resolution": "No earlier than 2027-07-27"},
        {"gate": "daily matured base predictions", "current": daily_eval["sample"]["matured_base_predictions"], "required": daily_contract["prospective_protocol"]["minimum_matured_base_predictions"], "status": "FAIL_NOT_MATURE", "earliest_resolution": "After 252 valid post-registration sessions"},
        {"gate": "daily matured selective signals", "current": daily_eval["sample"]["matured_selective_signals"], "required": daily_contract["prospective_protocol"]["minimum_matured_selective_signals"], "status": "FAIL_NOT_MATURE", "earliest_resolution": "Path dependent; no backfill"},
        {"gate": "weekly committed origins", "current": weekly_eval["sample"]["committed_weeks"], "required": weekly_eval["sample"]["minimum_committed_weeks"], "status": "FAIL_NOT_MATURE", "earliest_resolution": "Approximately 2028-W30 if no misses"},
        {"gate": "external generalization", "current": "DA 45.42%; lift -0.42 pp", "required": "DA/BDA >=55%; lift CI lower >0", "status": "FAILED_LOCKED_TEST", "earliest_resolution": "Cannot be repaired by retuning the closed family"},
        {"gate": "executable cost/fill validation", "current": False, "required": True, "status": "MISSING", "earliest_resolution": "Requires preregistered timestamped bid/ask and fill policy"},
        {"gate": "independent model validation", "current": False, "required": True, "status": "MISSING", "earliest_resolution": "After sufficient prospective evidence"},
        {"gate": "capital authorization", "current": False, "required": "separate governed decision", "status": "PROHIBITED", "earliest_resolution": "Never automatic"},
    ])

    transport_metrics = pd.read_csv(ROOT / INPUTS["transport_metrics"])
    transport_annual = pd.read_csv(ROOT / INPUTS["transport_annual"])
    provenance = pd.DataFrame([
        {
            "input_id": key,
            "relative_path": relative,
            "bytes": (ROOT / relative).stat().st_size,
            "sha256": sha256(ROOT / relative),
        }
        for key, relative in INPUTS.items()
    ])
    frames = {
        "Claim_Ladder": claim_ladder,
        "Model_Status": model_status,
        "Research_Metrics": research_metrics,
        "Failed_Families": failed_families,
        "Acceptance_Gaps": gaps,
        "Transport_Metrics": transport_metrics,
        "Transport_Annual": transport_annual,
        "Provenance": provenance,
    }
    summary = {
        "overall_status": "NOT_INSTITUTIONAL_DIRECTIONAL_CLAIM_READY",
        "highest_supported_claim_level": "research_candidate",
        "current_directional_trials": transport_registration["directional_trials_after"],
        "current_global_trials": transport_registration["global_trials_after"],
        "weekly_committed": weekly_eval["sample"]["committed_weeks"],
        "daily_committed": daily_eval["sample"]["committed_sessions"],
        "external_transport_pass": transport["primary_pass"],
        "signal_authorized": False,
        "capital_authorized": False,
        "truthful_current_claim": (
            "One exploratory H1 candidate is operationally ready for prospective adjudication; "
            "directional skill and external generalization are not yet demonstrated."
        ),
    }
    return frames, summary


def write_workbook(frames: dict[str, pd.DataFrame], summary: dict[str, Any]) -> None:
    readme = pd.DataFrame([
        {"item": "Overall status", "detail": summary["overall_status"]},
        {"item": "Highest supported claim", "detail": summary["highest_supported_claim_level"]},
        {"item": "Truthful current claim", "detail": summary["truthful_current_claim"]},
        {"item": "Directional/global trials", "detail": f"{summary['current_directional_trials']} / {summary['current_global_trials']}"},
        {"item": "Prospective observations", "detail": f"weekly={summary['weekly_committed']}; daily={summary['daily_committed']}"},
        {"item": "External transport", "detail": "FAIL" if not summary["external_transport_pass"] else "PASS"},
        {"item": "Capital", "detail": "NOT AUTHORIZED"},
        {"item": "Method", "detail": "Read-only packaging of sealed evidence; no fit, prediction, selection or new trial."},
    ])
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        readme.to_excel(writer, index=False, sheet_name="README")
        for sheet, frame in frames.items():
            frame.to_excel(writer, index=False, sheet_name=sheet)

    workbook = load_workbook(WORKBOOK)
    header_fill = PatternFill("solid", fgColor="17324D")
    pass_fill = PatternFill("solid", fgColor="D9EAD3")
    fail_fill = PatternFill("solid", fgColor="F4CCCC")
    pending_fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            sheet.column_dimensions[column[0].column_letter].width = min(max(max(values) + 2, 12), 52)
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            text = " ".join(str(cell.value or "") for cell in row).upper()
            fill = fail_fill if any(token in text for token in ("FAIL", "REJECT", "NOT_REACHED", "PROHIBITED")) else None
            if fill is None and any(token in text for token in ("NOT_YET", "INSUFFICIENT", "MISSING")):
                fill = pending_fill
            if fill is None and "PASS" in text:
                fill = pass_fill
            if fill is not None:
                for cell in row:
                    cell.fill = fill
    for sheet_name in ("Research_Metrics", "Transport_Metrics", "Transport_Annual"):
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        percent_headers = {
            "coverage", "directional_accuracy", "balanced_accuracy", "causal_majority_accuracy",
            "lift", "lift_ci_low", "lift_ci_high", "up_recall", "down_recall",
            "minimum_class_recall", "prediction_up_rate", "actual_up_rate",
            "wilson_95_low", "wilson_95_high", "lift_vs_causal_majority",
        }
        for name in percent_headers.intersection(headers):
            for cells in sheet.iter_rows(min_row=2, min_col=headers[name], max_col=headers[name]):
                cells[0].number_format = "0.00%"
    workbook.save(WORKBOOK)


def readme_text(summary: dict[str, Any]) -> str:
    return f"""# USD/COP directional institutional due-diligence pack v1

Generated from sealed artifacts. This pack performs no model fitting, prediction,
selection or metric recomputation.

## Current verdict

`{summary['overall_status']}`

Highest supported claim: **research candidate**. The truthful statement today is:

> {summary['truthful_current_claim']}

The historical H1 candidate is interesting but not conclusive: 2020–2024 selective
DA was 55.79% on 95 signals, while paired lift had a 95% interval crossing zero and
p=0.052. The locked external USD/MXN + USD/BRL transport failed with 45.42% DA,
45.35% BDA and −0.42 pp lift. Weekly and daily prospective ledgers currently have
zero observations because their first eligible dates have not arrived.

## What would change the verdict

- Complete at least 12 months, 252 matured daily base predictions and 100 matured
  selective signals without changing the rule or backfilling misses.
- Pass DA, BDA, both recalls, class balance, paired lift, PT, Brier and time-uniform
  evidence gates.
- Add a separately preregistered executable quote/fill and all-in cost ledger.
- Obtain independent model-validation and governance sign-off.

No dashboard, generator, backtest or this workbook can authorize capital.

## Files

- `usdcop_directional_due_diligence_v1.xlsx`: human review workbook.
- `evidence_manifest.json`: machine-readable status plus input/output hashes.
- CSV mirrors of every evidence table for independent agents.

Trial accounting: **{summary['current_directional_trials']} directional / {summary['current_global_trials']} global**.
"""


def validate_outputs(expected_sheets: list[str], output_hashes: dict[str, str]) -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    expected = ["README", *expected_sheets]
    if workbook.sheetnames != expected:
        raise RuntimeError(f"Workbook sheet drift: {workbook.sheetnames} != {expected}")
    workbook.close()
    for relative, expected_hash in output_hashes.items():
        if sha256(OUTPUT_DIR / relative) != expected_hash:
            raise RuntimeError(f"Output hash drift: {relative}")


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    frames, summary = evidence_frames()
    write_workbook(frames, summary)
    for sheet, frame in frames.items():
        frame.to_csv(OUTPUT_DIR / f"{sheet.lower()}.csv", index=False)
    (OUTPUT_DIR / "README.md").write_text(readme_text(summary), encoding="utf-8")
    output_names = [
        WORKBOOK.name,
        "README.md",
        *[f"{sheet.lower()}.csv" for sheet in frames],
    ]
    output_hashes = {name: sha256(OUTPUT_DIR / name) for name in output_names}
    manifest = {
        "schema_version": "1.0.0",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        **summary,
        "exporter_sha256": sha256(Path(__file__)),
        "input_sha256": {
            key: {"path": relative, "sha256": sha256(ROOT / relative)}
            for key, relative in INPUTS.items()
        },
        "output_sha256": output_hashes,
        "no_model_fit_prediction_selection_or_new_trial": True,
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    validate_outputs(list(frames), output_hashes)
    print(json.dumps({
        "status": summary["overall_status"],
        "highest_supported_claim_level": summary["highest_supported_claim_level"],
        "workbook": str(WORKBOOK.relative_to(ROOT)),
        "sheets": ["README", *frames],
        "inputs_hashed": len(INPUTS),
        "outputs_hashed": len(output_hashes),
        "directional_trials": summary["current_directional_trials"],
        "global_trials": summary["current_global_trials"],
        "capital_authorized": False,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
