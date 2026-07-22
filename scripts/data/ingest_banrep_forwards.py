#!/usr/bin/env python3
"""
Ingest BanRep USD/COP forward market monthly series (PIT) — TAREA C1
====================================================================

Point-in-time monthly ingestion of the official Banco de la República
USD/COP forward market series (montos negociados + devaluación implícita
anualizada, by tenor bucket) into ``macro_banrep_forwards_monthly``.
0 trials: data only, no study is opened here.

Source (documented 2026-07-21)
------------------------------
- Human catalog: SUAMECA > Catálogo > Sector externo, tasas de cambio y
  derivados > Mercado de derivados > "Otras series del Mercado de forwards"
  (menu id 430502 in
  https://suameca.banrep.gov.co/estadisticas-economicas-back/rest/estadisticaEconomicaRestService/consultaMenuXopcion?opcion=CATALOGO_DATOS,
  dashboard id 4160203 in ...?opcion=DASHBOARD).
- Machine endpoint (direct Excel, no auth):
  https://suameca.banrep.gov.co/archivos/sector_externo_tasas_cambio_derivados/mercado_derivados/series_historico_otros_derivados.xlsx
  Sheet "2. FwdUSDCOP": Fecha (month) | Reportante | Contraparte | Rango |
  MontoNegociado (USD mn) | DevaluacionImplicita (annualized, decimal).
- The BanRep SDMX web service (totoro.banrep.gov.co/nsi-jax-ws/rest/data)
  only exposes IBR/DTF/TRM/TPM/TIB/COLCAP/M-aggregates/UVR — no forwards —
  so the Excel above is the authoritative machine-readable endpoint.

Coverage honesty (task premise said 1997->)
-------------------------------------------
BanRep labels this series "Disponible desde 2005"; the file holds 2005-01
onward, complete. The 1997-2004 monthly forward series existed only on the
legacy OBIEE portal (totoro.banrep.gov.co/analytics), which now answers
"Sitio en Mantenimiento" — it is not machine-readable anywhere on the
current BanRep web. We load the full official series: 2005-01 -> present.

Publication lag (defines the PIT date)
--------------------------------------
Observed 2026-07-21: latest month in the file = 2026-05 (present at
month_end + 51 calendar days) while 2026-06 is absent (month_end + 21d):
the file refreshes roughly mid-M+2. Repo precedent for the sibling PDF
bulletin (config/usdcop_forward_macro_sources.yaml,
banrep_monthly_derivatives) is observation + 45 calendar days. We store the
more conservative
    published_at = last calendar day of month M + 60 days
so a PIT join on published_at <= as_of over-waits rather than under-waits.

Usage
-----
    set -a; . ./.env; set +a            # load POSTGRES_* credentials
    python scripts/data/ingest_banrep_forwards.py                # full ingest
    python scripts/data/ingest_banrep_forwards.py --dry-run      # no DB writes
    python scripts/data/ingest_banrep_forwards.py --from-xlsx f.xlsx  # offline

Applies migration database/migrations/065_banrep_forwards_monthly.sql
automatically if the table does not exist (psql-style: executes the SQL file).
"""

from __future__ import annotations

import argparse
import calendar
import os
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]  # scripts/data/<this> -> repo root
MIGRATION_FILE = PROJECT_ROOT / "database" / "migrations" / "065_banrep_forwards_monthly.sql"

XLSX_URL = (
    "https://suameca.banrep.gov.co/archivos/sector_externo_tasas_cambio_derivados/"
    "mercado_derivados/series_historico_otros_derivados.xlsx"
)
SHEET_NAME = "2. FwdUSDCOP"
CACHE_DIR = PROJECT_ROOT / "data" / "cache" / "banrep_forwards"  # gitignored, regenerable
SOURCE_TAG = "banrep_suameca_otros_derivados_xlsx"
PUBLICATION_LAG_DAYS = 60  # conservative: observed ~51d worst case, precedent 45d

TENOR_BUCKETS = [
    "0", "1 a 3", "4 a 14", "15 a 35", "36 a 60", "61 a 90", "91 a 180", "mayor a 180",
]

# Gaps documented as legitimate (none known as of 2026-07: the series is
# complete 2005-01 onward). Add (year, month) tuples if BanRep ever skips one.
DOCUMENTED_MISSING_MONTHS: set[tuple[int, int]] = set()

EXPECTED_START = date(2005, 1, 1)   # cobertura oficial documentada de la serie
PUBLICATION_LAG_DAYS = 60           # regla PIT: published_at = fin de mes + 60d


# ---------------------------------------------------------------------------
# Env / DB helpers
# ---------------------------------------------------------------------------

def _load_dotenv_fallback() -> None:
    """Populate POSTGRES_* from repo .env if not already in the environment."""
    if os.environ.get("POSTGRES_PASSWORD"):
        return
    env_path = PROJECT_ROOT / ".env"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key, value = key.strip(), value.strip().strip('"').strip("'")
        if key.startswith("POSTGRES") and key not in os.environ:
            os.environ[key] = value


def get_connection():
    import psycopg2

    return psycopg2.connect(
        host=os.environ.get("POSTGRES_HOST", "localhost"),
        port=int(os.environ.get("POSTGRES_PORT", "5432")),
        dbname=os.environ.get("POSTGRES_DB", "usdcop_trading"),
        user=os.environ.get("POSTGRES_USER", "admin"),
        password=os.environ.get("POSTGRES_PASSWORD", ""),
    )


def ensure_table(conn) -> None:
    """Apply migration 065 if macro_banrep_forwards_monthly does not exist."""
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_schema = 'public'
                  AND table_name = 'macro_banrep_forwards_monthly'
            )
            """
        )
        exists = cur.fetchone()[0]
    if exists:
        return
    print(f"[MIGRATE] applying {MIGRATION_FILE.name}")
    sql = MIGRATION_FILE.read_text(encoding="utf-8")
    with conn.cursor() as cur:
        cur.execute(sql)
    conn.commit()
    print("[MIGRATE] macro_banrep_forwards_monthly created")


# ---------------------------------------------------------------------------
# Download + parse
# ---------------------------------------------------------------------------

def download_xlsx(retries: int = 3) -> Path:
    """Download the official Excel to the (gitignored) cache dir."""
    import requests

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    dest = CACHE_DIR / "series_historico_otros_derivados.xlsx"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    last_err: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            resp = requests.get(XLSX_URL, headers=headers, timeout=180)
            resp.raise_for_status()
            if len(resp.content) < 100_000:
                raise ValueError(f"Suspiciously small download ({len(resp.content)} bytes)")
            fd, tmpname = tempfile.mkstemp(dir=CACHE_DIR, suffix=".part")
            os.close(fd)  # Windows: keep no open handle or replace() fails (WinError 32)
            tmp = Path(tmpname)
            tmp.write_bytes(resp.content)
            tmp.replace(dest)
            print(f"[DOWNLOAD] {dest} ({dest.stat().st_size:,} bytes)")
            return dest
        except Exception as exc:  # noqa: BLE001 — retry then surface
            last_err = exc
            print(f"[WARN] download attempt {attempt}/{retries} failed: {exc}")
    raise RuntimeError(
        f"Could not download {XLSX_URL}: {last_err}\n"
        f"BLOCKED fallback: open https://suameca.banrep.gov.co/estadisticas-economicas/catalogo"
        f" > Sector externo... > Mercado de derivados > 'Otras series del Mercado de forwards',"
        f" download the Excel manually and re-run with --from-xlsx <file>."
    )


def _month_of(value) -> date:
    """Normalize the sheet's Fecha cell (datetime, first of month) to a month DATE."""
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    # string fallback e.g. '2005-01-01'
    d = date.fromisoformat(str(value)[:10])
    return date(d.year, d.month, 1)


def published_at_for(month: date) -> date:
    """Conservative PIT date: last day of month + PUBLICATION_LAG_DAYS."""
    month_end = date(month.year, month.month, calendar.monthrange(month.year, month.month)[1])
    return month_end + timedelta(days=PUBLICATION_LAG_DAYS)


def parse_xlsx(path: Path) -> list[dict]:
    """Parse sheet '2. FwdUSDCOP' -> aggregated (month, tenor) rows.

    Raw grain is (month, Reportante=IMC, Contraparte, Rango). We aggregate to
    (month, tenor): monto = sum(MontoNegociado); implied_dev = monto-weighted
    mean of DevaluacionImplicita over rows where it is present. BanRep's own
    'Total' rows (Contraparte='Total', Rango='Total') are ingested verbatim
    as tenor='TOTAL'.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"[FAIL] sheet {SHEET_NAME!r} not found; sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    header_seen = False
    # (month, tenor) -> [sum_monto, sum_monto_dev_weighted, sum_weight_with_dev]
    agg: dict[tuple[date, str], list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0])
    totals: dict[date, tuple[float, float | None]] = {}

    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            if row and row[0] == "Fecha":
                expected = ("Fecha", "Reportante", "Contraparte", "Rango",
                            "MontoNegociado", "DevaluacionImplicita")
                got = tuple(str(c) for c in row[:6])
                if got != expected:
                    raise SystemExit(f"[FAIL] header drift: {got} != {expected}")
                header_seen = True
            continue
        if not row or row[0] is None:
            continue
        month = _month_of(row[0])
        rango = str(row[3]).strip()
        monto = float(row[4]) if row[4] is not None else 0.0
        dev = float(row[5]) if row[5] is not None else None
        # validacion de fila FUENTE antes de agregar (Codex verify C1 #5): un valor
        # invalido no puede cancelarse/promediarse dentro de un agregado "valido"
        if monto < 0:
            raise SystemExit(f"[FAIL] fila fuente con monto negativo: {month} {rango!r} {monto}")
        if dev is not None and not (-1.0 < dev < 2.0):
            raise SystemExit(f"[FAIL] fila fuente con dev implausible: {month} {rango!r} {dev}")

        if rango == "Total":
            totals[month] = (monto, dev)
            continue
        if rango not in TENOR_BUCKETS:
            raise SystemExit(f"[FAIL] unknown Rango bucket {rango!r} at {month} — "
                             f"update TENOR_BUCKETS after verifying the source")
        acc = agg[(month, rango)]
        acc[0] += monto
        if dev is not None:
            acc[1] += monto * dev
            acc[2] += monto
    wb.close()

    if not header_seen:
        raise SystemExit(f"[FAIL] header row not found in sheet {SHEET_NAME!r}")

    rows: list[dict] = []
    for (month, tenor), (monto, wsum, w) in agg.items():
        rows.append({
            "month": month,
            "tenor": tenor,
            "forward_rate": None,  # not published in this series (see migration 065)
            "implied_dev": (wsum / w) if w > 0 else None,
            "monto_negociado_usd_mn": monto,
            "published_at": published_at_for(month),
            "source": SOURCE_TAG,
        })
    for month, (monto, dev) in totals.items():
        rows.append({
            "month": month,
            "tenor": "TOTAL",
            "forward_rate": None,
            "implied_dev": dev,
            "monto_negociado_usd_mn": monto,
            "published_at": published_at_for(month),
            "source": SOURCE_TAG,
        })
    rows.sort(key=lambda r: (r["month"], r["tenor"]))
    return rows


# ---------------------------------------------------------------------------
# Validation (hard-fails only on data integrity)
# ---------------------------------------------------------------------------

def _month_iter(start: date, end: date):
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        yield date(y, m, 1)
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)


def validate(rows: list[dict]) -> None:
    if not rows:
        raise SystemExit("[FAIL] no rows parsed")

    # 1) key uniqueness
    keys = [(r["month"], r["tenor"]) for r in rows]
    if len(keys) != len(set(keys)):
        raise SystemExit("[FAIL] duplicated (month, tenor) keys in parsed data")

    # 2) value sanity: montos >= 0, dev inside (-1, 2)
    bad_monto = [r for r in rows if r["monto_negociado_usd_mn"] is not None
                 and r["monto_negociado_usd_mn"] < 0]
    if bad_monto:
        raise SystemExit(f"[FAIL] negative montos, e.g. {bad_monto[:3]}")
    bad_dev = [r for r in rows if r["implied_dev"] is not None
               and not (-1.0 < r["implied_dev"] < 2.0)]
    if bad_dev:
        raise SystemExit(f"[FAIL] implausible implied_dev, e.g. {bad_dev[:3]}")

    # 3) huecos sobre el SPAN ESPERADO completo (Codex verify C1 #2/#3): desde
    #    EXPECTED_START hasta el ultimo mes exigible segun el rezago PIT — un
    #    archivo truncado pero continuo tambien FALLA; cualquier hueco no
    #    documentado es FAIL duro, sin tolerancia.
    months = sorted({r["month"] for r in rows})
    have = set(months)
    today = date.today()
    last_exigible = months[0]
    for mo in _month_iter(EXPECTED_START, date(today.year, today.month, 1)):
        if published_at_for(mo) <= today:
            last_exigible = max(last_exigible, mo)
    missing = [
        mo for mo in _month_iter(EXPECTED_START, max(last_exigible, months[-1]))
        if mo not in have and (mo.year, mo.month) not in DOCUMENTED_MISSING_MONTHS
    ]
    if missing:
        raise SystemExit(
            f"[FAIL] {len(missing)} undocumented missing months in expected span "
            f"{EXPECTED_START} -> {max(last_exigible, months[-1])} (first 12: "
            f"{missing[:12]}). If legitimate, add to DOCUMENTED_MISSING_MONTHS."
        )
    print(f"[OK] no missing months in expected span {EXPECTED_START} -> "
          f"{max(last_exigible, months[-1])}")

    # 4) TOTAL cross-check BLOQUEANTE (Codex verify C1 #4): cada mes debe traer la
    #    fila TOTAL de BanRep y buckets comparables; divergencia >= 0.01 = FAIL.
    by_month: dict[date, dict[str, dict]] = defaultdict(dict)
    for r in rows:
        by_month[r["month"]][r["tenor"]] = r
    worst, n_cmp, skipped = 0.0, 0, []
    for month, tenors in sorted(by_month.items()):
        tot = tenors.get("TOTAL")
        buckets = [t for k, t in tenors.items() if k != "TOTAL"]
        if not tot or tot["implied_dev"] is None or not buckets:
            skipped.append(month)
            continue
        w = sum(b["monto_negociado_usd_mn"] or 0.0 for b in buckets
                if b["implied_dev"] is not None)
        if w <= 0:
            skipped.append(month)
            continue
        ours = sum((b["monto_negociado_usd_mn"] or 0.0) * b["implied_dev"]
                   for b in buckets if b["implied_dev"] is not None) / w
        diff = abs(ours - tot["implied_dev"])
        n_cmp += 1
        worst = max(worst, diff)
    if skipped:
        raise SystemExit(f"[FAIL] {len(skipped)} months without comparable TOTAL/"
                         f"buckets for cross-check (first: {skipped[:6]})")
    if worst >= 0.01:
        raise SystemExit(f"[FAIL] bucket-weighted dev vs BanRep TOTAL diverges: "
                         f"worst |diff| = {worst:.4f} >= 0.01")
    print(f"[XCHECK OK] bucket-weighted dev vs BanRep TOTAL: {n_cmp} months, "
          f"worst |diff| = {worst:.4f}")

    print(f"[OK] validation passed: {len(rows)} rows, {months[0]} -> {months[-1]}")


def report_by_decade(rows: list[dict]) -> None:
    decades: dict[int, set] = defaultdict(set)
    tenor_counts: dict[int, int] = defaultdict(int)
    for r in rows:
        dec = r["month"].year // 10 * 10
        decades[dec].add(r["month"])
        tenor_counts[dec] += 1
    print("[REPORT] coverage per decade:")
    for dec in sorted(decades):
        print(f"    {dec}s : {len(decades[dec])} months, {tenor_counts[dec]} (month,tenor) rows")


# ---------------------------------------------------------------------------
# Upsert
# ---------------------------------------------------------------------------

UPSERT_SQL = """
INSERT INTO macro_banrep_forwards_monthly
    (month, tenor, forward_rate, implied_dev, monto_negociado_usd_mn,
     published_at, source)
VALUES
    (%(month)s, %(tenor)s, %(forward_rate)s, %(implied_dev)s,
     %(monto_negociado_usd_mn)s, %(published_at)s, %(source)s)
ON CONFLICT (month, tenor) DO UPDATE SET
    forward_rate           = EXCLUDED.forward_rate,
    implied_dev            = EXCLUDED.implied_dev,
    monto_negociado_usd_mn = EXCLUDED.monto_negociado_usd_mn,
    published_at           = EXCLUDED.published_at,
    source                 = EXCLUDED.source,
    ingested_at            = now()
"""


def upsert(conn, rows: list[dict]) -> None:
    from psycopg2.extras import execute_batch

    with conn.cursor() as cur:
        execute_batch(cur, UPSERT_SQL, rows, page_size=500)
    conn.commit()
    with conn.cursor() as cur:
        cur.execute(
            "SELECT count(*), count(DISTINCT month), min(month), max(month) "
            "FROM macro_banrep_forwards_monthly"
        )
        n, nm, lo, hi = cur.fetchone()
    print(f"[DB] macro_banrep_forwards_monthly now holds {n} rows "
          f"({nm} months), {lo} -> {hi}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Ingest BanRep USD/COP forward market monthly series (PIT)."
    )
    parser.add_argument("--dry-run", action="store_true",
                        help="download+parse+validate only, no DB writes")
    parser.add_argument("--from-xlsx", default=None,
                        help="offline mode: parse a manually downloaded Excel")
    args = parser.parse_args()

    print("=" * 70)
    print("BanRep forwards USD/COP mensual -> macro_banrep_forwards_monthly (TAREA C1)")
    print(f"  sheet   : {SHEET_NAME} of series_historico_otros_derivados.xlsx")
    print(f"  source  : {XLSX_URL}")
    print(f"  PIT     : published_at = month_end + {PUBLICATION_LAG_DAYS}d (conservative)")
    print("=" * 70)

    if args.from_xlsx:
        path = Path(args.from_xlsx)
        print(f"[OFFLINE] parsing {path}")
    else:
        path = download_xlsx()

    rows = parse_xlsx(path)
    validate(rows)
    report_by_decade(rows)

    if args.dry_run:
        print("[DRY-RUN] skipping DB writes")
        return 0

    _load_dotenv_fallback()
    conn = get_connection()
    try:
        ensure_table(conn)
        upsert(conn, rows)
    finally:
        conn.close()

    print("[DONE] ingest complete (0 trials: data only, no study opened)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
