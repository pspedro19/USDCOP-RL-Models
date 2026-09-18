"""Build editable and PDF thesis documents from pinned retrospective evidence.

No training, live calls, policy search, secret loading or hold-out replay. New output
directory required. Markdown is the writing source; figures and tables are computed.
"""

# ruff: noqa: RUF001
# Mathematical multiplication signs in Spanish figure labels are intentional.
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import importlib.metadata
import json
import re
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
DEPS = ROOT / "outputs/thesis-delivery/_deps"
if DEPS.exists():
    sys.path.append(str(DEPS))

SOURCE = ROOT / "outputs/thesis-repair/research_grade_20260912_v5"
SOURCE_SHA = "00fc276e22b62ffac144bca1abb6c468f15e272225c46743304c67fe120ae2da"
SEEDS = [42, 123, 456, 789, 1337]
LABELS = {
    "always_flat": "No operar",
    "NULL_A_short": "Siempre corto 1×",
    "B1_session": "Largo intradía 1×",
    "B1_passive_uncosted": "Pasivo sin costos*",
    "ppo_regime_mean5": "PPO régimen: cartera 5",
    "ppo_backbone_mean5": "PPO sin régimen: cartera 5",
    "ppo_median_weights": "PPO pesos medianos",
    "LogReg": "Regresión logística",
    "deepseek": "LLM DeepSeek",
    "azure": "LLM Azure",
    "hybrid_deepseek": "Híbrido DeepSeek",
    "hybrid_azure": "Híbrido Azure",
}
PANEL_ARMS = [
    "always_flat",
    "ppo_median_weights",
    "deepseek",
    "azure",
    "hybrid_deepseek",
    "hybrid_azure",
]
COLORS = ["#4b5563", "#0072B2", "#D55E00", "#CC79A7", "#009E73", "#6f63a5"]
LINESTYLES = [":", "-", "--", "-.", "-", "--"]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_json(path):
    def reject(value):
        raise ValueError(f"nonfinite JSON: {value}")

    return json.loads(Path(path).read_text(encoding="utf-8"), parse_constant=reject)


def write_json(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, indent=2, ensure_ascii=False, allow_nan=False)


def compound(values):
    x = np.asarray(values, dtype=float)
    if x.ndim != 1 or not len(x) or not np.isfinite(x).all() or np.any(x <= -1):
        raise ValueError("finite daily returns greater than -1 required")
    return float(np.expm1(np.log1p(x).sum()) * 100)


def curve(values):
    compound(values)
    return 100 * np.r_[1.0, np.cumprod(1 + np.asarray(values))]


def break_even(gross, cost):
    """Unique root within [0,1], NOT an assumed executable spread."""
    g, c = np.asarray(gross, float), np.asarray(cost, float)
    if g.shape != c.shape or g.ndim != 1 or not np.isfinite([g, c]).all() or (c < 0).any():
        raise ValueError("aligned finite gross and nonnegative costs required")
    if compound(g) < 0 or not np.any(c > 0):
        return None
    if compound(g) == 0:
        return 0.0
    if compound(g - c) > 0:
        return None
    low, high = 0.0, 1.0
    for _ in range(70):
        mid = (low + high) / 2
        if compound(g - mid * c) > 0:
            low = mid
        else:
            high = mid
    return (low + high) / 2


def verify_source(source=SOURCE):
    if sha(source / "manifest.json") != SOURCE_SHA:
        raise ValueError("source manifest pin changed")
    manifest = load_json(source / "manifest.json")
    for name, expected in manifest["artifacts_sha256"].items():
        path = (source / name).resolve()
        if path.parent != source.resolve() or sha(path) != expected:
            raise ValueError("source artifact mismatch: " + name)
    # Authenticate every consumed archived input without unpickling or retraining.
    from scripts.presentation.build_research_grade_thesis import Snapshot

    snapshot = Snapshot(manifest["snapshot_manifest"])
    if sha(snapshot.path) != manifest["snapshot_sha256"]:
        raise ValueError("snapshot manifest mismatch")
    for name, expected in manifest["inputs_sha256"].items():
        if hashlib.sha256(snapshot.read(name)).hexdigest() != expected:
            raise ValueError("archived source mismatch: " + name)
    report, daily = load_json(source / "results.json"), load_json(source / "daily_series.json")
    dates = [r["date"] for r in daily["always_flat"]]
    if len(dates) != 226 or dates != sorted(set(dates)):
        raise ValueError("invalid cohort")
    for name, rows in daily.items():
        if [r["date"] for r in rows] != dates:
            raise ValueError("unaligned cohort: " + name)
        g, c, n = (
            np.array([r[key] for r in rows], float)
            for key in ("gross_return", "cost_return", "net_return")
        )
        if (
            not np.isfinite([g, c, n]).all()
            or (c < 0).any()
            or not np.allclose(g - c, n, atol=1e-12, rtol=0)
        ):
            raise ValueError("daily accounting: " + name)
        m = report["metrics"][name]
        eq = curve(n)
        checks = {
            "return_compounded_pct": compound(n),
            "gross_compounded_pct": compound(g),
            "gross_sum_pct": g.sum() * 100,
            "cost_sum_pct": c.sum() * 100,
            "max_drawdown_pct": (eq / np.maximum.accumulate(eq) - 1).min() * 100,
        }
        for key, value in checks.items():
            if not np.isclose(value, m[key], atol=1e-9, rtol=0):
                raise ValueError(f"metric mismatch {name}/{key}")
    return report, daily, snapshot


def diagnostics(report, daily):
    from src.research.reporting_v2 import bootstrap_indices

    protocol = load_json(ROOT / "docs/thesis/diagnostic_protocol.json")
    p = protocol["bootstrap"]
    if protocol["source_manifest_sha256"] != SOURCE_SHA:
        raise ValueError("diagnostic protocol/source mismatch")
    rows = daily["ppo_median_weights"]
    g = np.array([r["gross_return"] for r in rows])
    idx = bootstrap_indices(
        len(g), p["replications"], p["seed"], tuple(p["expected_block_lengths"])
    )
    summed = g[idx].sum(axis=1) * 100
    compounded = np.expm1(np.log1p(g[idx]).sum(axis=1)) * 100
    sum_ci = np.percentile(summed, [2.5, 97.5])
    if not np.allclose(sum_ci, report["ppo_median_gross_sum_ci95_pct"], atol=1e-9, rtol=0):
        raise ValueError("published gross interval not reproduced")
    result = {
        "scope": "retrospective_descriptive",
        "confirmatory": False,
        "alpha_established": False,
        "bootstrap": p,
        "gross_sum_ci95_pct": sum_ci.tolist(),
        "gross_compounded_ci95_pct": np.percentile(compounded, [2.5, 97.5]).tolist(),
        "break_even_cost_multiplier": {},
        "quarters": [],
        "limitations": [
            "conditional on frozen policy",
            "not selection-adjusted",
            "not a risk-factor alpha",
            "assumed costs",
        ],
    }
    for name in report["headline_order"]:
        r = daily[name]
        result["break_even_cost_multiplier"][name] = break_even(
            [v["gross_return"] for v in r], [v["cost_return"] for v in r]
        )
    for q in range(1, 5):
        r = [v for v in rows if (int(v["date"][5:7]) - 1) // 3 + 1 == q]
        result["quarters"].append(
            {
                "quarter": q,
                "n_sessions": len(r),
                "gross_compounded_pct": compound([v["gross_return"] for v in r]),
                "gross_sum_pct": sum(v["gross_return"] for v in r) * 100,
                "net_compounded_pct": compound([v["net_return"] for v in r]),
            }
        )
    return result


def write_csv(path, rows):
    with Path(path).open("x", newline="", encoding="utf-8-sig") as h:
        writer = csv.DictWriter(h, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def figures(out, report, daily, extra):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    import pandas as pd

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "svg.hashsalt": "thesis-retrospective-delivery",
            "savefig.dpi": 240,
        }
    )
    target = out / "figures"
    target.mkdir()
    dates = pd.to_datetime([r["date"] for r in daily["always_flat"]])
    # Initial value is a plotting anchor, not an invented trading session.
    x = dates.insert(0, dates[0] - pd.Timedelta(days=1))
    actions = load_json(SOURCE / "actions_regime.json")
    regime = load_json(SOURCE / "regime_provenance.json")
    k_regime, counts = regime["declared_k"], regime["session_counts"]
    if len(counts) != k_regime or sum(counts) != len(dates):
        raise ValueError("regime counts do not match cohort")
    for arm in {r["arm"] for r in actions}:
        for k in range(k_regime):
            group = [r for r in actions if r["arm"] == arm and r["regime_id"] == k]
            if (
                len(group) != 3
                or {r["action"] for r in group} != {"SHORT", "FLAT", "LONG"}
                or sum(r["count"] for r in group) != counts[k] * 59
                or any(r["total"] != counts[k] * 59 for r in group)
                or any(not np.isclose(r["pct"], 100 * r["count"] / r["total"]) for r in group)
            ):
                raise ValueError("invalid action/regime denominator")
    m = report["metrics"]

    def draw(ax, panel, compact=False):
        if panel in (0, 1, 2):
            for name, color, style in zip(PANEL_ARMS, COLORS, LINESTYLES, strict=True):
                net = np.array([r["net_return"] for r in daily[name]])
                eq = curve(net)
                if panel == 0:
                    ax.plot(x, eq, label=LABELS[name], color=color, ls=style, lw=1.55)
                elif panel == 1:
                    ax.plot(
                        x,
                        (eq / np.maximum.accumulate(eq) - 1) * 100,
                        label=LABELS[name],
                        color=color,
                        ls=style,
                        lw=1.55,
                    )
                elif name != "always_flat":
                    sr = pd.Series(net).rolling(60, min_periods=60)
                    y = sr.mean() / sr.std(ddof=1).replace(0, np.nan) * np.sqrt(221)
                    counts_window = (
                        pd.Series([r["round_trips_or_lower_bound"] for r in daily[name]])
                        .rolling(60, min_periods=60)
                        .sum()
                    )
                    y = y.where(counts_window >= 20)
                    ax.plot(dates, y, label=LABELS[name], color=color, ls=style, lw=1.55)
            ax.xaxis.set_major_locator(mdates.MonthLocator(interval=3))
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%Y"))
            ax.set_ylabel(
                [
                    "Capital normalizado (inicio = 100)",
                    "Caída desde el máximo (%)",
                    "Sharpe (reloj de 221 sesiones)",
                ][panel]
            )
            if panel == 2:
                ax.axhline(0, color="#777777", lw=0.7)
            if not compact or panel == 0:
                ax.legend(fontsize=7 if compact else 9, ncol=2, loc="best")
        elif panel == 3:
            for name, color, style in zip(PANEL_ARMS[1:], COLORS[1:], LINESTYLES[1:], strict=True):
                g = np.array([r["gross_return"] for r in daily[name]])
                c = np.array([r["cost_return"] for r in daily[name]])
                ax.plot(
                    [0, 1, 2, 3],
                    [compound(g - k * c) for k in [0, 1, 2, 3]],
                    color=color,
                    ls=style,
                    marker="o",
                    label=LABELS[name],
                )
            ax.axhline(0, color="#777777", lw=0.7)
            ax.set_xticks([0, 1, 2, 3], ["0 (bruto)", "×1", "×2", "×3"])
            ax.set_ylabel("Retorno compuesto (%)")
            ax.set_xlabel("Multiplicador del costo supuesto · posición fija")
            if not compact:
                ax.legend(fontsize=8, ncol=2)
        elif panel == 4:
            bottom = np.zeros(k_regime)
            for action, color, hatch in [
                ("SHORT", "#0072B2", "//"),
                ("FLAT", "#b8bec7", ""),
                ("LONG", "#009E73", ".."),
            ]:
                vals = [
                    next(
                        r["pct"]
                        for r in actions
                        if r["arm"] == "hybrid_deepseek"
                        and r["regime_id"] == k
                        and r["action"] == action
                    )
                    for k in range(k_regime)
                ]
                ax.bar(
                    range(k_regime),
                    vals,
                    bottom=bottom,
                    color=color,
                    hatch=hatch,
                    edgecolor="white",
                    label=action,
                )
                bottom += vals
            ax.set_xticks(range(k_regime), [f"R{k}\n{counts[k]} ses." for k in range(k_regime)])
            ax.set_ylim(0, 100)
            ax.set_ylabel("Decisiones del híbrido DeepSeek (%)")
            ax.set_xlabel("K = 5; quinta coordenada recuperada sólo para descripción")
            ax.legend(fontsize=8, ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.0))
        else:
            groups = [
                [m[f"{c}_seed{s}"]["sharpe"] for s in SEEDS] for c in ["ppo_regime", "ppo_backbone"]
            ]
            ax.boxplot(groups, widths=0.4)
            for j, vals in enumerate(groups):
                for k, (seed, value) in enumerate(zip(SEEDS, vals, strict=True)):
                    ax.scatter(
                        j + 1 + (k - 2) * 0.035, value, color=["#0072B2", "#009E73"][j], s=28
                    )
                    if not compact:
                        ax.annotate(
                            str(seed),
                            (j + 1 + (k - 2) * 0.035, value),
                            xytext=(6, 1),
                            textcoords="offset points",
                            fontsize=7,
                        )
            ax.set_xticks([1, 2], ["PPO régimen", "PPO sin régimen"])
            ax.set_ylabel("Sharpe neto por semilla")
            ax.axhline(0, color="#777777", lw=0.7)
            ax.set_xlabel("Cinco semillas por configuración; no son réplicas LLM")
        ax.set_title(
            [
                "A · Capital neto — selección 2023",
                "B · Drawdown comparativo",
                "C · Sharpe móvil — 60 sesiones",
                "D · Sensibilidad a costos",
                "E · Acciones por régimen (descriptivo)",
                "F · Variabilidad entre semillas",
            ][panel],
            fontsize=11,
            loc="left",
            pad=12,
        )
        ax.grid(axis="y", alpha=0.2)

    def save(fig, stem):
        for ext in ("png", "svg", "pdf"):
            fig.savefig(
                target / f"{stem}.{ext}",
                bbox_inches="tight",
                metadata={"Creator": "USDCOP thesis evidence renderer"} if ext == "pdf" else {},
            )
        plt.close(fig)

    stems = [
        "01_capital",
        "02_drawdown",
        "03_sharpe_movil",
        "04_costos",
        "05_acciones_regimen",
        "06_semillas",
    ]
    fig, axs = plt.subplots(2, 3, figsize=(18, 10.6), layout="constrained")
    for i, ax in enumerate(axs.ravel()):
        draw(ax, i, True)
    fig.suptitle(
        "USD/COP · evidencia retrospectiva real\nSelección 2023 · 226 sesiones · costos supuestos · no confirmatorio",
        fontsize=18,
    )
    save(fig, "00_resumen_seis_paneles")
    for i, stem in enumerate(stems):
        fig, ax = plt.subplots(figsize=(9.3, 5.4), layout="constrained")
        draw(ax, i)
        fig.supxlabel(
            "Fuente: bundle retrospectivo v5 · no representa operaciones ejecutadas en un venue",
            fontsize=8,
        )
        save(fig, stem)
    fig, axs = plt.subplots(1, 2, figsize=(12, 4.8), layout="constrained")
    names = ["NULL_A_short", "ppo_median_weights", "LogReg", "hybrid_deepseek", "deepseek", "azure"]
    pos = np.arange(len(names))
    axs[0].barh(
        pos - 0.18,
        [m[n]["gross_compounded_pct"] for n in names],
        0.35,
        label="Bruto compuesto",
        color="#0072B2",
    )
    axs[0].barh(
        pos + 0.18,
        [m[n]["return_compounded_pct"] for n in names],
        0.35,
        label="Neto compuesto",
        color="#D55E00",
    )
    axs[0].set_yticks(pos, [LABELS[n] for n in names])
    axs[0].axvline(0, color="gray", lw=0.7)
    axs[0].legend(fontsize=8)
    axs[0].set_title("Bruto y neto: misma convención")
    axs[0].set_xlabel("Retorno compuesto (%)")
    gross = [r["gross_return"] for r in daily["ppo_median_weights"]]
    net = [r["net_return"] for r in daily["ppo_median_weights"]]
    axs[1].plot(x, curve(gross), label="PPO mediano · bruto", color="#0072B2")
    axs[1].plot(x, curve(net), label="PPO mediano · neto", color="#D55E00", ls="--")
    axs[1].axhline(100, color="gray", ls=":")
    axs[1].set_ylabel("Capital normalizado")
    axs[1].legend(fontsize=8)
    axs[1].xaxis.set_major_locator(mdates.MonthLocator(interval=3))
    axs[1].xaxis.set_major_formatter(mdates.DateFormatter("%m/%Y"))
    axs[1].set_title("Dos trayectorias, no una resta de acumulados")
    save(fig, "07_bruto_neto")
    dq = load_json(SOURCE / "data_quality.json")["flat_ohlc_fraction_by_year"]
    fig, ax = plt.subplots(figsize=(9.3, 4.6), layout="constrained")
    ax.bar(list(dq), [100 * v for v in dq.values()], color="#0072B2")
    ax.set_ylabel("Barras con O = H = L = C (%)")
    ax.set_title(
        "Calidad de representación del histórico; no es un resultado de estrategia", loc="left"
    )
    ax.set_ylim(0, 108)
    for i, v in enumerate(dq.values()):
        ax.text(i, 100 * v + 1, f"{100 * v:.1f}%", ha="center", fontsize=9)
    save(fig, "08_calidad_datos")
    fig, ax = plt.subplots(figsize=(9.3, 4.2), layout="constrained")
    cis = [extra["gross_sum_ci95_pct"], extra["gross_compounded_ci95_pct"]]
    points = [
        m["ppo_median_weights"]["gross_sum_pct"],
        m["ppo_median_weights"]["gross_compounded_pct"],
    ]
    for i, (pt, ci) in enumerate(zip(points, cis, strict=True)):
        ax.errorbar(pt, i, xerr=[[pt - ci[0]], [ci[1] - pt]], fmt="o", capsize=7, color="#0072B2")
        ax.annotate(
            f"{pt:.2f}% · IC95 [{ci[0]:.2f}, {ci[1]:.2f}]",
            (pt, i),
            xytext=(0, 14),
            textcoords="offset points",
            ha="center",
        )
    ax.axvline(0, color="#D55E00", ls="--")
    ax.set_yticks([0, 1], ["Suma bruta", "Bruto compuesto"])
    ax.set_ylim(-0.6, 1.7)
    ax.set_xlabel("Porcentaje · bootstrap estacionario por sesiones, 10.000 réplicas")
    ax.set_title(
        "PPO mediano: resultado positivo observado, incertidumbre compatible con cero",
        loc="left",
        fontsize=11,
    )
    save(fig, "09_incertidumbre_bruto")
    fig, ax = plt.subplots(figsize=(11, 5.4), layout="constrained")
    ax.axis("off")
    nodes = [
        (0.13, 0.76, "Precios M5\nUSD/COP"),
        (0.13, 0.28, "Contexto macro\nidentidad / fecha / disponibilidad"),
        (0.42, 0.52, "Dataset congelado\nfeatures / scaler / régimen"),
        (0.70, 0.80, "Supervisado / PPO\nmodelos históricos"),
        (0.70, 0.52, "LLM DeepSeek / Azure\ndecisiones numéricas registradas"),
        (0.70, 0.23, "Híbrido\nacuerdo de dirección"),
        (0.92, 0.52, "Motor común\nposiciones\nbruto / costo / neto"),
    ]
    for x0, y0, label in nodes:
        ax.text(
            x0,
            y0,
            label,
            ha="center",
            va="center",
            fontsize=9,
            bbox={"boxstyle": "round,pad=.6", "facecolor": "#e7eff5", "edgecolor": "#7894aa"},
        )
    for start, end in [
        ((0.22, 0.76), (0.32, 0.56)),
        ((0.22, 0.28), (0.32, 0.48)),
        ((0.52, 0.55), (0.60, 0.78)),
        ((0.52, 0.52), (0.58, 0.52)),
        ((0.70, 0.72), (0.70, 0.32)),
        ((0.70, 0.44), (0.70, 0.32)),
        ((0.81, 0.80), (0.87, 0.58)),
        ((0.83, 0.52), (0.86, 0.52)),
        ((0.79, 0.23), (0.87, 0.45)),
    ]:
        ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "->", "color": "#52697c"})
    ax.set_xlim(0, 1.06)
    ax.set_ylim(0, 1)
    ax.set_title("Arquitectura de la evaluación retrospectiva", loc="left", fontsize=14)
    ax.text(
        0.5,
        0.04,
        "No hay noticias en los prompts de esta cohorte. Flechas = flujo lógico; no certifican disponibilidad histórica.",
        ha="center",
        fontsize=9,
    )
    save(fig, "10_arquitectura")


def fmt(value, digits=2):
    return "NA" if value is None else f"{value:.{digits}f}".replace(".", ",")


def table(headers, rows):
    return "\n".join(
        ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
        + ["| " + " | ".join(map(str, row)) + " |" for row in rows]
    )


def expand_content(text, report, extra):
    text = re.sub(
        r"^FIGURE: ([^|\n]+)\| (.+)$",
        lambda m: f"![{m[2].strip()}]({m[1].strip()})",
        text,
        flags=re.M,
    )
    m = report["metrics"]
    replacements = {
        "{{TABLE_GLOBAL}}": table(
            ["Brazo", "Bruto comp. %", "Neto comp. %", "Sharpe neto", "MaxDD %"],
            [
                [
                    LABELS[n],
                    fmt(m[n]["gross_compounded_pct"]),
                    fmt(m[n]["return_compounded_pct"]),
                    fmt(m[n]["sharpe"]),
                    fmt(m[n]["max_drawdown_pct"]),
                ]
                for n in report["headline_order"]
            ],
        ),
        "{{TABLE_SEEDS}}": table(
            ["Configuración", "Semilla", "Bruto comp. %", "Neto comp. %", "Sharpe"],
            [
                [
                    c.replace("ppo_", ""),
                    s,
                    fmt(m[f"{c}_seed{s}"]["gross_compounded_pct"]),
                    fmt(m[f"{c}_seed{s}"]["return_compounded_pct"]),
                    fmt(m[f"{c}_seed{s}"]["sharpe"]),
                ]
                for c in ["ppo_regime", "ppo_backbone"]
                for s in SEEDS
            ],
        ),
        "{{TABLE_COSTS}}": table(
            ["Política congelada", "Costo ×0 %", "Costo ×1 %", "Costo ×2 %", "Costo ×3 %"],
            [
                [LABELS[n]] + [fmt(v[k]) for k in ["zero_cost", "x1", "x2", "x3"]]
                for n, v in load_json(SOURCE / "cost_stress.json").items()
            ],
        ),
        "{{TABLE_BREAK_EVEN}}": table(
            ["Brazo", "Fracción del costo base en equilibrio"],
            [
                [LABELS[n], "Sin raíz en [0,1]" if v is None else fmt(v, 4)]
                for n, v in extra["break_even_cost_multiplier"].items()
                if n != "always_flat"
            ],
        ),
        "{{TABLE_QUARTERS}}": table(
            ["Trimestre 2023", "Sesiones", "Suma bruta %", "Bruto comp. %", "Neto comp. %"],
            [
                [
                    r["quarter"],
                    r["n_sessions"],
                    fmt(r["gross_sum_pct"]),
                    fmt(r["gross_compounded_pct"]),
                    fmt(r["net_compounded_pct"]),
                ]
                for r in extra["quarters"]
            ],
        ),
        "{{TABLE_STATISTICS}}": table(
            [
                "Brazo frente a flat",
                "Media diaria neta %",
                "IC95 de media %",
                "p centrado",
                "p Holm",
            ],
            [
                [
                    LABELS[k.removesuffix("_vs_flat")],
                    fmt(v["estimate"] * 100, 4),
                    "[" + "; ".join(fmt(x * 100, 4) for x in v["ci95"]) + "]",
                    fmt(v["p_centered_stationary"], 6),
                    fmt(v["p_holm_retrospective_family"], 6),
                ]
                for k, v in report["paired_primary_mean_return_tests"].items()
            ],
        ),
        "{{COMPOUND_CI}}": "["
        + "; ".join(fmt(v) for v in extra["gross_compounded_ci95_pct"])
        + "]",
        "{{BREAK_EVEN_MEDIAN}}": fmt(extra["break_even_cost_multiplier"]["ppo_median_weights"], 4),
        "{{COST_REDUCTION_MEDIAN}}": fmt(
            100 * (1 - extra["break_even_cost_multiplier"]["ppo_median_weights"])
        ),
    }
    for token, value in replacements.items():
        text = text.replace(token, value)
    if re.search(r"\{\{.+?\}\}", text):
        raise ValueError("unexpanded manuscript token")
    return text


def strip_frontmatter(text):
    return re.sub(r"\A---\r?\n.*?\r?\n---\r?\n", "", text, count=1, flags=re.S)


def blocks(text):
    lines, i = text.splitlines(), 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("#"):
            n = len(line) - len(line.lstrip("#"))
            yield "heading", (n, line[n:].strip())
            i += 1
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*)\]\(([^)]+)\)", line)
            if not match:
                raise ValueError("invalid figure declaration")
            yield "figure", (match[1], match[2])
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", x) for x in cells):
                    rows.append(cells)
                i += 1
            yield "table", rows
        else:
            para = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "![")):
                para.append(lines[i].strip())
                i += 1
            yield "paragraph", " ".join(para)


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text).replace("**", "").replace("`", "")
    return re.sub(r"\*([^*\n]+)\*", r"\1", text)


def render_document(text, stem, out, short=False):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from matplotlib import get_data_path
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    fontdir = Path(get_data_path()) / "fonts/ttf"
    pdfmetrics.registerFont(TTFont("Thesis", str(fontdir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("ThesisBold", str(fontdir / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily(
        "Thesis", normal="Thesis", bold="ThesisBold", italic="Thesis", boldItalic="ThesisBold"
    )
    styles = {
        "body": ParagraphStyle(
            "body",
            fontName="Thesis",
            fontSize=10,
            leading=15,
            spaceAfter=8,
            alignment=TA_JUSTIFY,
            allowWidows=0,
            allowOrphans=0,
        ),
        "caption": ParagraphStyle(
            "caption", fontName="Thesis", fontSize=8, leading=11, spaceAfter=12
        ),
        "cell": ParagraphStyle("cell", fontName="Thesis", fontSize=7.4, leading=10),
        "h1": ParagraphStyle(
            "h1", fontName="ThesisBold", fontSize=19, leading=25, spaceAfter=17, keepWithNext=True
        ),
        "h2": ParagraphStyle(
            "h2",
            fontName="ThesisBold",
            fontSize=12,
            leading=17,
            spaceBefore=13,
            spaceAfter=7,
            keepWithNext=True,
        ),
    }

    def para(value, style="body"):
        return Paragraph(html.escape(plain(value)), styles[style])

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(7)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string("17324D")
    footer = section.footer.paragraphs[0]
    footer.text = "USD/COP · memoria para revisión académica | "
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    title = "Sistema de trading algorítmico para USD/COP mediante aprendizaje automático supervisado y aprendizaje por refuerzo"
    subtitle = (
        "Capítulos 3 y 4 · evidencia retrospectiva"
        if short
        else "Memoria de maestría · versión integral para revisión del director"
    )
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle)
    doc.add_paragraph(
        "Esp. Ing. Pedro Elias Perez Salazar\nMaestría en Inteligencia Artificial · UBA-FIUBA\nDirector: Esp. Ing. Diego Asencio\n14 de septiembre de 2026"
    )
    doc.add_paragraph(
        "Resultados empíricos retrospectivos; no certifica rentabilidad operativa. Esta edición reemplaza las cifras sintéticas del documento de presentación, no los archivos originales."
    )
    story = [
        Spacer(1, 65),
        para(title, "h1"),
        Spacer(1, 20),
        para(subtitle),
        Spacer(1, 25),
        para("Esp. Ing. Pedro Elias Perez Salazar"),
        para("Maestría en Inteligencia Artificial · UBA-FIUBA"),
        para("Director: Esp. Ing. Diego Asencio"),
        para("14 de septiembre de 2026"),
        Spacer(1, 35),
        para(
            "Resultados empíricos retrospectivos. Versión para revisión académica; no certifica rentabilidad operativa."
        ),
    ]
    parsed = list(blocks(text))
    headings = [v[1] for k, v in parsed if k == "heading" and v[0] <= 2]
    doc.add_page_break()
    doc.add_heading("Contenido", 1)
    story += [PageBreak(), para("Contenido", "h1")]
    for h in headings:
        doc.add_paragraph(h)
        story.append(para(h, "caption"))
    for block_index, (kind, val) in enumerate(parsed):
        if kind == "heading":
            level, title = val
            if level == 1:
                doc.add_page_break()
                story.append(PageBreak())
            doc.add_heading(title, min(level, 3))
            story.append(para(title, "h1" if level == 1 else "h2"))
        elif kind == "paragraph":
            is_caption = val.startswith("Tabla ")
            doc.add_paragraph(plain(val), style="Caption" if is_caption else "Normal")
            story.append(para(val, "caption" if is_caption else "body"))
        elif kind == "table":
            t = doc.add_table(rows=1, cols=len(val[0]))
            t.style = "Light Shading Accent 1"
            for j, v in enumerate(val[0]):
                t.rows[0].cells[j].text = plain(v)
            for row in val[1:]:
                cells = t.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = plain(v)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8)
            header_repeat = OxmlElement("w:tblHeader")
            t.rows[0]._tr.get_or_add_trPr().append(header_repeat)
            widths = [465 / len(val[0])] * len(val[0])
            if len(val[0]) >= 4:
                widths = [145] + [320 / (len(val[0]) - 1)] * (len(val[0]) - 1)
            pt = Table(
                [[para(v, "cell") for v in row] for row in val],
                colWidths=widths,
                repeatRows=1,
                hAlign="LEFT",
            )
            pt.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7eff5")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#f5f7f9")],
                        ),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            pt.keepWithNext = True
            spacer = Spacer(1, 12)
            spacer.keepWithNext = True
            story += [pt, spacer]
        else:
            caption, relative = val
            path = (out / relative).resolve()
            path.relative_to(out.resolve())
            if not path.is_file():
                raise ValueError("figure missing: " + relative)
            # Give every chart a dedicated page so labels remain readable in Word.
            doc.add_page_break()
            doc.add_picture(str(path), width=Cm(16.4))
            doc.add_paragraph(caption, style="Caption")
            from PIL import Image as PILImage

            with PILImage.open(path) as im:
                ratio = im.height / im.width
            figure_group = [
                Image(str(path), width=465, height=465 * ratio),
                para(caption, "caption"),
            ]
            if (
                block_index > 0
                and parsed[block_index - 1][0] == "heading"
                and parsed[block_index - 1][1][0] > 1
            ):
                # Keep a section heading with its first figure without forcing an
                # otherwise unnecessary page after a split preceding paragraph.
                figure_group.insert(0, story.pop())
            story.append(PageBreak())
            story.append(KeepTogether(figure_group))
    doc.save(out / f"{stem}.docx")

    def header(canvas, pdfdoc):
        canvas.setFont("Thesis", 8)
        canvas.setFillColor(colors.HexColor("#596777"))
        canvas.drawString(65, 813, "USD/COP · evidencia retrospectiva · revisión académica")
        canvas.drawRightString(530, 32, str(pdfdoc.page))

    pdf = SimpleDocTemplate(
        str(out / f"{stem}.pdf"),
        pagesize=A4,
        leftMargin=65,
        rightMargin=65,
        topMargin=57,
        bottomMargin=53,
        title=subtitle,
        author="Pedro Elias Perez Salazar",
    )
    pdf.build(story, onFirstPage=header, onLaterPages=header)


def render_review(out):
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    from pypdf import PdfReader

    review = out / "review"
    review.mkdir()
    audit = {}
    for stem in ["tesis_completa", "capitulos_3_4"]:
        pdf = pdfium.PdfDocument(out / f"{stem}.pdf")
        extracted = PdfReader(out / f"{stem}.pdf")
        thumbs, pages = [], []
        for i, page in enumerate(pdf):
            im = page.render(scale=1.2).to_pil()
            im.save(review / f"{stem}_{i + 1:02}.png")
            thumb = im.copy()
            thumb.thumbnail((298, 425))
            thumbs.append(thumb)
            text = extracted.pages[i].extract_text() or ""
            pages.append(
                {
                    "page": i + 1,
                    "text_characters": len(text),
                    "has_image": bool(extracted.pages[i].images),
                    "replacement_character": "\ufffd" in text,
                }
            )
            page.close()
        for start in range(0, len(thumbs), 9):
            sheet = Image.new("RGB", (940, 1370), "#d6dde4")
            draw = ImageDraw.Draw(sheet)
            for j, thumb in enumerate(thumbs[start : start + 9]):
                x, y = 10 + (j % 3) * 310, 24 + (j // 3) * 455
                sheet.paste(thumb, (x, y))
                draw.text((x, y - 18), f"{stem} - {start + j + 1}", fill="black")
            sheet.save(review / f"{stem}_contact_{start // 9 + 1:02}.png")
        pdf.close()
        audit[stem] = {
            "page_count": len(pages),
            "pages": pages,
            "rendered_all_pages": True,
            "human_visual_review": "PENDING",
        }
    write_json(out / "document_qa.json", audit)


def build(output):
    out = Path(output).resolve()
    out.relative_to((ROOT / "outputs/thesis-delivery").resolve())
    if out.exists():
        raise FileExistsError("new output directory required; existing deliveries are immutable")
    report, daily, snapshot = verify_source()
    extra = diagnostics(report, daily)
    out.mkdir(parents=True)
    source_dir = out / "evidence"
    source_dir.mkdir()
    for name in [
        "manifest.json",
        "results.json",
        "metrics.csv",
        "daily_series.json",
        "actions_regime.json",
        "cost_stress.json",
        "data_quality.json",
        "regime_provenance.json",
    ]:
        shutil.copyfile(SOURCE / name, source_dir / name)
    for name in [
        "regime5_20260914_preparation_post_review.json",
        "usdcop_daily_cross_source_v2.json",
    ]:
        shutil.copyfile(ROOT / "outputs/thesis-repair" / name, source_dir / name)
    shutil.copyfile(
        ROOT / "docs/thesis/diagnostic_protocol.json", source_dir / "diagnostic_protocol.json"
    )
    write_json(out / "diagnostics.json", extra)
    write_csv(
        out / "metrics_full.csv",
        [{"arm": name, **report["metrics"][name]} for name in report["headline_order"]],
    )
    write_csv(
        out / "seed_metrics.csv",
        [
            {"configuration": c, "seed": s, **report["metrics"][f"{c}_seed{s}"]}
            for c in ["ppo_regime", "ppo_backbone"]
            for s in SEEDS
        ],
    )
    write_csv(out / "quarter_diagnostics.csv", extra["quarters"])
    # Numeric matrix behind all figures; no hand-drawn or simulated equity.
    write_csv(
        out / "daily_plot_data.csv",
        [{"arm": a, **row} for a, rows in daily.items() for row in rows],
    )
    figures(out, report, daily, extra)
    source_text = strip_frontmatter(
        (ROOT / "docs/thesis/chapters_1_2_5.md").read_text(encoding="utf-8")
    )
    c34 = strip_frontmatter((ROOT / "docs/thesis/chapters_3_4.md").read_text(encoding="utf-8"))
    before, after = source_text.split("# 5.", 1)
    main34, appendix = c34.split("# Anexo técnico", 1)
    chapter5, bibliography = after.split("# Bibliografía", 1)
    full = expand_content(
        before
        + "\n"
        + main34
        + "\n# 5."
        + chapter5
        + "\n# Anexo técnico"
        + appendix
        + "\n# Bibliografía"
        + bibliography,
        report,
        extra,
    )
    short = expand_content(c34, report, extra)
    # The chapter excerpt includes references too; never a dangling author-year citation.
    if "# Bibliografía" in source_text:
        short += "\n# Bibliografía" + source_text.split("# Bibliografía", 1)[1]
    for stem, text in [("tesis_completa", full), ("capitulos_3_4", short)]:
        (out / f"{stem}.md").write_text(text, encoding="utf-8")
        render_document(text, stem, out, stem == "capitulos_3_4")
    claims = [
        {
            "claim": "global_returns",
            "source": "evidence/results.json:metrics",
            "section": "4.2",
            "status": "reproduced",
        },
        {
            "claim": "ppo_median_10_percent",
            "source": "evidence/daily_series.json:ppo_median_weights",
            "section": "4.3",
            "status": "observed_gross_not_alpha",
        },
        {
            "claim": "compound_gross_uncertainty",
            "source": "diagnostics.json:gross_compounded_ci95_pct",
            "section": "4.3",
            "status": "retrospective_conditional",
        },
        {
            "claim": "cost_break_even",
            "source": "diagnostics.json:break_even_cost_multiplier",
            "section": "4.5",
            "status": "fixed_policy_counterfactual",
        },
        {
            "claim": "regime_actions",
            "source": "evidence/regime_provenance.json",
            "section": "4.7",
            "status": "descriptive_reclassification_only",
        },
        {
            "claim": "llm_complete_cohort",
            "source": "evidence/results.json:ledger_quality",
            "section": "3.6",
            "status": "not_original_request_certification",
        },
        {
            "claim": "future_38_training",
            "source": "evidence/regime5_20260914_preparation_post_review.json",
            "section": "3.9",
            "status": "not_executed",
        },
        {
            "claim": "seed_dispersion",
            "source": "seed_metrics.csv",
            "section": "4.4",
            "status": "five_seeds_per_configuration",
        },
        {
            "claim": "quarters",
            "source": "quarter_diagnostics.csv",
            "section": "4.4",
            "status": "all_four_quarters_not_selection",
        },
        {
            "claim": "h2_uncertainty",
            "source": "evidence/results.json:h2_hierarchical",
            "section": "4.4",
            "status": "ci_includes_zero",
        },
        {
            "claim": "ohlc_representation",
            "source": "evidence/data_quality.json",
            "section": "4.7",
            "status": "descriptive_not_executability",
        },
        {
            "claim": "daily_cross_source",
            "source": "evidence/usdcop_daily_cross_source_v2.json",
            "section": "3.2",
            "status": "prior_diagnostic_not_rerun_network",
        },
    ]
    for name in sorted((out / "figures").glob("*.png")):
        claims.append(
            {
                "claim": "figure_" + name.stem,
                "source": str(name.relative_to(out)),
                "section": "4",
                "status": "derived_from_pinned_inputs_and_diagnostics",
            }
        )
    write_csv(out / "claim_evidence_matrix.csv", claims)
    render_review(out)
    inputs = {
        str(p.relative_to(ROOT)): sha(p)
        for p in [
            Path(__file__),
            ROOT / "docs/thesis/chapters_1_2_5.md",
            ROOT / "docs/thesis/chapters_3_4.md",
            ROOT / "docs/thesis/diagnostic_protocol.json",
            ROOT / "src/research/reporting_v2.py",
            ROOT / "src/research/inference.py",
            ROOT / "scripts/presentation/build_research_grade_thesis.py",
        ]
    }
    original_refs = {}
    for name in [
        "memorianueva (2).pdf",
        "Capitulos_3-4_USDCOP_version_final_presentacion_sintetica.docx",
    ]:
        path = Path("C:/Users/pedro/Downloads") / name
        if path.exists():
            original_refs[name] = sha(path)
    manifest = {
        "contract": "THESIS-MANUSCRIPT-DELIVERY-1",
        "created_at_utc": datetime.now(UTC).isoformat(),
        "source_manifest_sha256": SOURCE_SHA,
        "source_inputs_verified": len(snapshot.used),
        "source_code_and_writing_sha256": inputs,
        "original_manuscripts_sha256": original_refs,
        "confirmatory": False,
        "profitability_established": False,
        "alpha_established": False,
        "no_training_or_provider_calls": True,
        "academic_review_required": True,
        "dependencies": {
            p: importlib.metadata.version(p)
            for p in [
                "numpy",
                "matplotlib",
                "pandas",
                "python-docx",
                "reportlab",
                "pypdfium2",
                "pypdf",
            ]
        },
        "artifacts_sha256": {
            str(p.relative_to(out)): sha(p) for p in sorted(out.rglob("*")) if p.is_file()
        },
    }
    write_json(out / "delivery_manifest.json", manifest)
    print(
        json.dumps(
            {
                "output": str(out),
                "documents": 4,
                "figures": len(list((out / "figures").glob("*.png"))),
                "gross_compounded_ci95": extra["gross_compounded_ci95_pct"],
                "break_even": extra["break_even_cost_multiplier"]["ppo_median_weights"],
                "source_verified": True,
            },
            ensure_ascii=False,
        )
    )
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    build(parser.parse_args().output)
