"""Excel hand-off for the causal USD/COP weekly directional replay."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADER_FILL = PatternFill("solid", fgColor="12304A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
UP_FILL = PatternFill("solid", fgColor="C6EFCE")
DOWN_FILL = PatternFill("solid", fgColor="FFC7CE")
FLAT_FILL = PatternFill("solid", fgColor="E7E6E6")
PASS_FILL = PatternFill("solid", fgColor="A9D18E")
FAIL_FILL = PatternFill("solid", fgColor="F4B084")
PERCENT_COLUMNS = {
    "coverage",
    "directional_accuracy",
    "balanced_accuracy",
    "up_recall",
    "down_recall",
    "minimum_class_recall",
    "prediction_up_rate",
    "actual_up_rate",
    "probability_up",
    "threshold",
    "evidence_da",
    "evidence_balanced_da",
    "evidence_min_recall",
    "evidence_score",
    "decision_confidence_proxy",
    "confidence_proxy",
    "validation_score",
    "forecast_log_return",
    "actual_log_return",
    "point_mae_log_return",
    "point_rmse_log_return",
    "point_naive_mae_log_return",
    "point_naive_rmse_log_return",
    "point_mae_skill_vs_spot",
    "point_rmse_skill_vs_spot",
    "point_mape_price",
    "point_directional_accuracy",
    "point_balanced_accuracy",
    "point_up_recall",
    "point_down_recall",
    "point_minimum_class_recall",
    "point_interval_coverage",
    "majority_baseline_da",
    "da_lift_vs_majority",
    "point_validation_mae_skill",
    "point_interval_level",
}
PERCENT_POINT_COLUMNS = {
    "forecast_return_pct", "point_abs_error_pct", "decision_forecast_return_pct",
}
PRICE_COLUMNS = {
    "base_price", "forecast_price", "forecast_price_change", "forecast_interval_lower",
    "forecast_interval_upper", "actual_price", "point_abs_error_price",
    "decision_forecast_price", "decision_forecast_interval_lower",
    "decision_forecast_interval_upper", "point_mae_price",
}


def _guide_frame(document: dict[str, Any]) -> pd.DataFrame:
    summaries = {int(item["year"]): item for item in document["summaries"]}
    y2026 = summaries.get(2026, {}).get("decision_metrics", {})
    warning = (
        "El DA seleccionado 2026 puede parecer alto, pero Balanced DA="
        f"{y2026.get('balanced_accuracy', 0):.1%} y recall mínimo="
        f"{y2026.get('minimum_class_recall', 0):.1%}; no tratarlo como edge promocionado."
    )
    rows = [
        ("asset", document["symbol"], "Instrumento; UP significa que sube USD/COP."),
        ("contract_hash", document["contract_hash"], "Identificador inmutable de esta publicación."),
        ("data_cutoff", document["data_cutoff"], "Último dato observable incluido."),
        ("latest_week", document["latest_week"], "Última inferencia semanal publicada."),
        ("status", document["methodology"]["status"], "RESEARCH_SHADOW: replay, no orden ejecutable."),
        ("signal_authorized", False, "Debe permanecer FALSE hasta promoción independiente."),
        (
            "regla_consumo",
            "Usar UP/DOWN solo si promotion_gate_passed=TRUE; en cualquier otro caso usar FLAT.",
            "Regla conservadora para otro agente o estrategia.",
        ),
        ("UP", "LONG USD / SHORT COP", "Dirección de USD/COP, no recomendación de tamaño."),
        ("DOWN", "SHORT USD / LONG COP", "Dirección de USD/COP, no recomendación de tamaño."),
        ("FLAT", "NO TRADE", "Abstención por desacuerdo, evidencia insuficiente o gate fallido."),
        (
            "selected_horizons",
            "Mejor candidato causal tactical + mejor candidato causal swing en cada origen.",
            "La selección solo usa outcomes cuyo target_date ya maduró.",
        ),
        ("H1", "timing", "Solo timing; no decide por sí solo la dirección multi-sleeve."),
        ("H5_H10_H15", "tactical", "Uno de estos horizontes representa el sleeve táctico."),
        ("H20_H25_H30", "swing", "Uno de estos horizontes representa el sleeve swing."),
        (
            "2025_training",
            "frozen_pre_2025",
            "Features/modelo fijados antes de 2025; ninguna etiqueta de 2025 entra al fit.",
        ),
        (
            "2026_training",
            "weekly_expanding_matured_labels",
            "Reentrena cada semana incorporando únicamente etiquetas ya maduras.",
        ),
        ("actual", "1=UP, 0=DOWN, vacío=pending", "Outcome realizado al target_date."),
        ("hit", "TRUE/FALSE/vacío", "Acierto únicamente cuando el outcome ya está disponible."),
        ("probability_up", "0..1", "Probabilidad del clasificador; comparar con threshold fijado."),
        (
            "forecast_price",
            "spot * exp(forecast_log_return)",
            "Punto forward del Ridge causal; es una estimación, no un precio garantizado.",
        ),
        (
            "forecast_interval",
            "intervalo residual 80% calibrado en 2022-2024",
            "Rango de incertidumbre alrededor del punto, no intervalo de ejecución.",
        ),
        (
            "direction_price_agree",
            "TRUE cuando clasificador y retorno continuo tienen el mismo signo",
            "La discrepancia es una alerta diagnóstica; no se oculta ni fuerza.",
        ),
        ("evidence_score", "score causal contraído", "Combina DA, Balanced DA y recall mínimo."),
        ("advertencia_metricas", warning, "Evita aprobar un clasificador de clase mayoritaria."),
        (
            "paso_1",
            "Filtrar WEEKLY_DECISIONS por la semana/origin_date que el agente conocía.",
            "Nunca consultar una fila futura durante replay.",
        ),
        (
            "paso_2",
            "Comprobar origin_is_partial_week y promotion_gate_passed.",
            "Una semana parcial puede cambiar cuando llegue el cierre semanal.",
        ),
        (
            "paso_3",
            "Leer decision_direction y selected_horizons; aplicar riesgo/tamaño fuera de este archivo.",
            "El contrato entrega dirección, no sizing ni autorización de ejecución.",
        ),
    ]
    return pd.DataFrame(rows, columns=["campo", "valor", "interpretacion"])


def _weekly_decisions_frame(document: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for week in document["weeks"]:
        decision = week["decision"]
        selected = [item for item in week["horizons"] if item["selected"]]
        rows.append({
            "iso_week": week["iso_week"],
            "year": week["year"],
            "origin_date": week["origin_date"],
            "origin_is_partial_week": week["origin_is_partial_week"],
            "base_price": week["base_price"],
            "regime": week["regime"]["state"],
            "direction_shift_z": week["regime"]["direction_shift_z"],
            "training_mode": week["training_mode"],
            "research_candidate_direction": decision["direction"],
            "gated_direction": (
                decision["direction"] if decision["promotion_gate_passed"] else "FLAT"
            ),
            "execution_direction": (
                decision["direction"]
                if decision["promotion_gate_passed"] and decision["signal_authorized"]
                else "FLAT"
            ),
            "decision_direction": decision["direction"],
            "decision_action": decision["action"],
            "decision_status": decision["status"],
            "signal_authorized": decision["signal_authorized"],
            "promotion_gate_passed": decision["promotion_gate_passed"],
            "primary_horizon": decision["primary_horizon"],
            "confirmation_horizons": "|".join(map(str, decision["confirmation_horizons"])),
            "selected_horizons": "|".join(map(str, decision["selected_horizons"])),
            "selected_predictions": "|".join(
                f"H{item['horizon_days']}:{item['prediction']}" for item in selected
            ),
            "selected_evidence_scores": "|".join(
                f"H{item['horizon_days']}:{item['evidence']['shrunk_score']:.4f}"
                for item in selected
            ),
            "confidence_proxy": decision["confidence_proxy"],
            "target_date": decision["target_date"],
            "forecast_price": decision["forecast_price"],
            "forecast_return_pct": decision["forecast_return_pct"],
            "forecast_interval_lower": decision["forecast_interval_lower"],
            "forecast_interval_upper": decision["forecast_interval_upper"],
            "actual": decision["actual"],
            "hit": decision["hit"],
            "rationale": decision["rationale"],
            "image_path": week["image_path"],
        })
    return pd.DataFrame(rows)


def _metrics_frame(document: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for summary in document["summaries"]:
        common = {
            "year": summary["year"],
            "weeks_total": summary["weeks_total"],
            "shadow_decisions": summary["shadow_decisions"],
            "abstentions": summary["abstentions"],
            "coverage": summary["coverage"],
        }
        rows.append({**common, "scope": "selected_strategy", "horizon_days": None,
                     **summary["decision_metrics"]})
        for metrics in summary["horizon_metrics"]:
            point = metrics.get("point_forecast", {})
            direction = {key: value for key, value in metrics.items() if key != "point_forecast"}
            rows.append({
                **common,
                "scope": "horizon",
                **direction,
                **{f"point_{key}": value for key, value in point.items()},
            })
    return pd.DataFrame(rows)


def _da_horizon_frame(document: dict[str, Any]) -> pd.DataFrame:
    """One auditable row per year/horizon, including the honest majority baseline."""
    rows: list[dict[str, Any]] = []
    by_horizon: dict[int, list[dict[str, Any]]] = {}
    for summary in document["summaries"]:
        for metrics in summary["horizon_metrics"]:
            actual_up = metrics["actual_up_rate"]
            majority_da = None if actual_up is None else max(actual_up, 1.0 - actual_up)
            da = metrics["directional_accuracy"]
            da_lift = None if da is None or majority_da is None else da - majority_da
            balanced = metrics["balanced_accuracy"] or 0.0
            minimum_recall = metrics["minimum_class_recall"] or 0.0
            if da_lift is not None and da_lift > 0 and balanced >= 0.52 and minimum_recall >= 0.30:
                verdict = "BUENO"
            elif balanced >= 0.52 and minimum_recall >= 0.30:
                verdict = "PROMETEDOR_BALANCEADO"
            elif balanced >= 0.50 and minimum_recall >= 0.20:
                verdict = "MIXTO"
            else:
                verdict = "DÉBIL_O_SESGADO"
            point = metrics["point_forecast"]
            row = {
                "year": summary["year"],
                "horizon_days": metrics["horizon_days"],
                "n": metrics["n"],
                "directional_accuracy": da,
                "balanced_accuracy": metrics["balanced_accuracy"],
                "up_recall": metrics["up_recall"],
                "down_recall": metrics["down_recall"],
                "minimum_class_recall": metrics["minimum_class_recall"],
                "prediction_up_rate": metrics["prediction_up_rate"],
                "actual_up_rate": actual_up,
                "majority_baseline_da": majority_da,
                "da_lift_vs_majority": da_lift,
                "brier": metrics["brier"],
                "verdict": verdict,
                "point_directional_accuracy": point["directional_accuracy"],
                "point_balanced_accuracy": point["balanced_accuracy"],
                "point_mae_skill_vs_spot": point["mae_skill_vs_spot"],
                "point_mae_price": point["mae_price"],
                "point_mape_price": point["mape_price"],
            }
            rows.append(row)
            by_horizon.setdefault(int(metrics["horizon_days"]), []).append(row)
    generalizes = {
        horizon: all(row["verdict"] == "BUENO" for row in values) and len(values) == 2
        for horizon, values in by_horizon.items()
    }
    for row in rows:
        row["generalizes_2025_2026"] = generalizes[int(row["horizon_days"])]
    return pd.DataFrame(rows)


def _features_frame(document: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for horizon_key, card in sorted(document["models"].items(), key=lambda item: int(item[0])):
        metrics = card["validation_metrics"]
        point = card["point_forecast"]
        point_metrics = point["validation_metrics"]
        for rank, feature in enumerate(card["selected_features"], start=1):
            rows.append({
                "horizon_days": int(horizon_key),
                "feature_rank": rank,
                "feature_name": feature,
                "model_family": card["model_family"],
                "half_life": card["half_life"],
                "threshold": card["threshold"],
                "validation_years": "|".join(map(str, card["validation_years"])),
                "validation_score": card["validation_score"],
                "validation_eligible": card["validation_eligible"],
                "validation_da": metrics["directional_accuracy"],
                "validation_balanced_da": metrics["balanced_accuracy"],
                "validation_min_recall": metrics["minimum_class_recall"],
                "point_model_family": point["model_family"],
                "point_alpha": point["alpha"],
                "point_validation_eligible": point["validation_eligible"],
                "point_validation_mae_skill": point_metrics["mae_skill_vs_spot"],
                "point_validation_mae_price": point_metrics["mae_price"],
                "point_interval_level": point["interval_level"],
                "feature_hash": card["feature_hash"],
            })
    return pd.DataFrame(rows)


def _lineage_frame(document: dict[str, Any]) -> pd.DataFrame:
    rows = [
        {"section": "contract", "key": "schema_version", "value": document["schema_version"]},
        {"section": "contract", "key": "contract_hash", "value": document["contract_hash"]},
        {"section": "contract", "key": "generated_at", "value": document["generated_at"]},
    ]
    rows.extend(
        {"section": "methodology", "key": key, "value": "|".join(map(str, value)) if isinstance(value, list) else value}
        for key, value in document["methodology"].items()
    )
    for key, source in document["lineage"].items():
        rows.append({"section": "lineage", "key": key, "value": source["path"]})
        rows.append({"section": "lineage_sha256", "key": key, "value": source["sha256"]})
    return pd.DataFrame(rows)


def _format_sheet(worksheet, table_name: str) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 30
    if worksheet.max_row > 1 and worksheet.max_column > 0:
        table = Table(displayName=table_name, ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        worksheet.add_table(table)

    headers = {cell.value: cell.column for cell in worksheet[1]}
    for name, column in headers.items():
        values = [str(name or "")]
        values.extend(str(worksheet.cell(row, column).value or "") for row in range(2, worksheet.max_row + 1))
        width = min(60, max(10, max(len(value) for value in values) + 2))
        worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = width
        if name in PERCENT_COLUMNS:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, column).number_format = "0.00%"
        if name in PERCENT_POINT_COLUMNS:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, column).number_format = '0.00"%"'
        if name in PRICE_COLUMNS:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, column).number_format = '#,##0.00'
        if name in {"rationale", "interpretacion", "valor"}:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)

    for direction_name in (
        "research_candidate_direction", "gated_direction", "execution_direction",
        "decision_direction", "prediction", "point_forecast_direction",
    ):
        column = headers.get(direction_name)
        if column:
            letter = worksheet.cell(1, column).column_letter
            cell_range = f"{letter}2:{letter}{worksheet.max_row}"
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"UP"'], fill=UP_FILL))
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"DOWN"'], fill=DOWN_FILL))
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"FLAT"'], fill=FLAT_FILL))
    gate_column = headers.get("promotion_gate_passed")
    if gate_column:
        letter = worksheet.cell(1, gate_column).column_letter
        cell_range = f"{letter}2:{letter}{worksheet.max_row}"
        worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=["TRUE"], fill=PASS_FILL))
        worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=["FALSE"], fill=FAIL_FILL))


def export_directional_replay_workbook(
    path: Path,
    document: dict[str, Any],
    ledger: pd.DataFrame,
    summary: pd.DataFrame,
) -> None:
    """Write one atomic, styled workbook suitable for human or agent hand-off."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.{os.getpid()}.tmp.xlsx")
    frames = {
        "GUIDE": _guide_frame(document),
        "WEEKLY_DECISIONS": _weekly_decisions_frame(document),
        "HORIZON_WEEKLY": ledger,
        "DA_HORIZON": _da_horizon_frame(document),
        "METRICS": _metrics_frame(document),
        "MODEL_FEATURES": _features_frame(document),
        "LINEAGE": _lineage_frame(document),
    }
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for index, (sheet_name, frame) in enumerate(frames.items(), start=1):
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.book[sheet_name]
                worksheet.sheet_properties.tabColor = "2DD4A8" if index == 2 else "12304A"
                _format_sheet(worksheet, f"ReplayTable{index}")
            writer.book.properties.title = "USD/COP weekly directional replay 2025-2026"
            writer.book.properties.subject = "Causal shadow forecasting hand-off"
            writer.book.properties.creator = "USDCOP Forecasting Pipeline"
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
