"""Contract and causality checks for the published USD/COP directional replay."""
from __future__ import annotations

import json
from datetime import date
import math
from pathlib import Path

from openpyxl import load_workbook
import pytest

from services.inference_api.contracts.forecasting import DirectionalReplayIndex
from src.forecasting.directional_replay import validate_replay_document


ROOT = Path(__file__).resolve().parents[2]
INDEX = (
    ROOT
    / "usdcop-trading-dashboard/public/forecasting/usdcop/directional_replay_index.json"
)


def _document() -> dict:
    assert INDEX.exists(), "run generate_usdcop_directional_replay.py first"
    return json.loads(INDEX.read_text(encoding="utf-8"))


def test_directional_replay_matches_backend_contract_and_images() -> None:
    document = _document()
    if hasattr(DirectionalReplayIndex, "model_validate"):
        DirectionalReplayIndex.model_validate(document)
    else:  # Pydantic v1 compatibility
        DirectionalReplayIndex.parse_obj(document)
    assert validate_replay_document(document, ROOT, require_images=True) == []


def test_directional_replay_has_every_week_and_every_horizon() -> None:
    document = _document()
    weeks = document["weeks"]
    latest_year, latest_week = map(
        int, document["latest_week"].replace("W", "").split("-")
    )
    expected = sum(
        date(year, 12, 28).isocalendar().week
        for year in range(2025, latest_year)
    ) + latest_week
    assert len(weeks) == expected
    assert weeks[0]["iso_week"] == "2025-W01"
    assert weeks[51]["iso_week"] == "2025-W52"
    assert weeks[52]["iso_week"] == "2026-W01"
    assert weeks[-1]["iso_week"] == document["latest_week"]
    assert all(
        [item["horizon_days"] for item in week["horizons"]]
        == [1, 5, 10, 15, 20, 25, 30]
        for week in weeks
    )


def test_2025_is_frozen_and_2026_uses_only_mature_labels() -> None:
    document = _document()
    for week in document["weeks"]:
        for horizon in week["horizons"]:
            label_end = horizon["train_label_end"]
            point_label_end = horizon["point_train_label_end"]
            assert label_end is None or label_end <= week["origin_date"]
            assert point_label_end is None or point_label_end <= week["origin_date"]
            if week["year"] == 2025:
                assert horizon["training_mode"] == "frozen_pre_2025"
                assert label_end is None or label_end < "2025-01-01"
                assert point_label_end is None or point_label_end < "2025-01-01"
            else:
                assert horizon["training_mode"] == "weekly_expanding_matured_labels"


def test_shadow_decision_never_claims_execution_authorization() -> None:
    for week in _document()["weeks"]:
        decision = week["decision"]
        assert decision["signal_authorized"] is False
        selected = [item for item in week["horizons"] if item["selected"]]
        assert sorted(item["horizon_days"] for item in selected) == sorted(
            decision["selected_horizons"]
        )
        if decision["direction"] != "FLAT":
            assert len(selected) == 2
            assert len({item["prediction"] for item in selected}) == 1
        if decision["promotion_gate_passed"]:
            assert decision["direction"] != "FLAT"
            assert all(item["eligible_for_direction"] for item in selected)


def test_forward_price_points_are_complete_and_numerically_consistent() -> None:
    document = _document()
    for week in document["weeks"]:
        for horizon in week["horizons"]:
            implied = week["base_price"] * math.exp(horizon["forecast_log_return"])
            assert horizon["forecast_price"] == pytest.approx(implied, abs=0.02)
            assert 0 < horizon["forecast_interval_lower"] <= horizon["forecast_price"]
            assert horizon["forecast_price"] <= horizon["forecast_interval_upper"]
            assert horizon["forecast_interval_level"] == pytest.approx(0.8)
            assert horizon["direction_price_agree"] is (
                horizon["prediction"] == horizon["point_forecast_direction"]
            )
            if horizon["actual"] is None:
                assert horizon["actual_price"] is None
                assert horizon["point_abs_error_price"] is None
    assert all(
        "point_forecast" in metrics
        for summary in document["summaries"]
        for metrics in summary["horizon_metrics"]
    )


def test_consolidated_excel_is_complete_and_agent_readable() -> None:
    document = _document()
    workbook_path = ROOT / "reports/usdcop_directional_replay_2025_2026.xlsx"
    assert workbook_path.exists()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "GUIDE", "WEEKLY_DECISIONS", "HORIZON_WEEKLY", "DA_HORIZON",
        "METRICS", "MODEL_FEATURES", "LINEAGE"
    ]
    assert workbook["WEEKLY_DECISIONS"].max_row == len(document["weeks"]) + 1
    assert workbook["HORIZON_WEEKLY"].max_row == len(document["weeks"]) * 7 + 1
    assert workbook["DA_HORIZON"].max_row == 15
    workbook.close()
